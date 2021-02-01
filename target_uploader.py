from openpyxl import load_workbook
import parser
from ptdb_request import PTDBRequest

loc = './assets/XOLO_PTDB_Nomenclature.xlsx'
wb = load_workbook(loc)
ws = wb.active
max_row = ws.max_row

for row in range(2, 3):
    target = {
        "target": ws.cell(row, 1).value,
        "partner": "",
        "protein_class_pk": "1",
        "notes": ws.cell(row, 2).value,
        "project": "Xolo",
        "subunits": parser.SubunitParser(ws, row).subunits,
    }
    PTDBRequest(target).post_target()

